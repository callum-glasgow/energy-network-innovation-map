"""Lossless ENA spreadsheet imports and GUID-preserving downloads for map updates."""
import copy
import csv
import io
import json
import re
from datetime import date, datetime
from pathlib import Path
from adapters.ena.client import DESCRIPTIVE, METADATA, apparent_test, catalogue, export_rows, reconcile, request
from core.config import OWNERS
from core.project_metadata import amount, iso_date, values
from core.storage import now
from core.text import clean, link

FIELD_MAP = json.loads(Path(__file__).with_name('export_fields.json').read_text())
ALIASES = {'id': ('id', 'Key', 'GUID', 'Project ID', 'Project GUID'),
           'projectTitle': ('title', 'Name'), 'projectReferenceNumber': ('reference', 'Ref'),
           'ownerNetwork': ('owner', 'OwnerNetwork'), 'projectStatus': ('status',),
           'projectStartDate': ('start_date', 'StartDate'), 'projectEndDate': ('end_date', 'EndDate'),
           'projectBudget': ('budget_raw',), 'strategyTheme': ('themes',),
           'technologyAreas': ('technology',), 'researchAreas': ('research',),
           'fundingMechanism': ('funding',), 'url': ('url', 'Project URL')}
MAX_ROWS = 20000

def header(value):
    return re.sub(r'[^a-z0-9]', '', str(value).casefold())


def comparable_value(key, value):
    if key in ('technologyAreas', 'strategyTheme', 'researchAreas', 'fundingMechanism'):
        return values(value)
    if key in ('projectStartDate', 'projectEndDate'):
        return iso_date(value) or clean(value)
    if key in ('projectBudget', 'projectPhaseCost', 'indicativeTotalNIAProjectExpenditure'):
        parsed = amount(value)
        return parsed if parsed is not None else clean(value)
    return clean(value)


def comparable_record(record):
    normalized = {key: comparable_value(key, value) for key, value in canonical(record).items()}
    return {key: value for key, value in normalized.items() if value is not None and value != '' and value != []}


def read_rows(path):
    path = Path(path)
    if path.suffix.lower() == '.json':
        data = json.loads(path.read_text(encoding='utf-8-sig'))
        rows = data.get('projects') if isinstance(data, dict) else data
    elif path.suffix.lower() == '.csv':
        raw = path.read_bytes()
        text = raw.decode('utf-16' if raw.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig')
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        if not reader.fieldnames or len(reader.fieldnames) != len(set(reader.fieldnames)):
            raise ValueError('Spreadsheet headers are missing or duplicated')
        rows = list(reader)
    elif path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if len(workbook.worksheets) != 1:
                raise ValueError('Use an ENA workbook with one sheet, or save the project sheet as CSV')
            sheet = workbook.worksheets[0]
            if sheet.max_row > MAX_ROWS + 1 or sheet.max_column > 256:
                raise ValueError('Spreadsheet exceeds the supported row/column limit')
            entries = sheet.iter_rows(values_only=True)
            names = [str(v).strip() if v is not None else '' for v in next(entries)]
            if not all(names) or len(names) != len(set(names)):
                raise ValueError('Spreadsheet headers are missing or duplicated')
            rows = [dict(zip(names, [v.isoformat() if isinstance(v, (date, datetime)) else v for v in row]))
                    for row in entries if any(v is not None for v in row)]
        finally:
            workbook.close()
    else:
        raise ValueError('Use an ENA .csv, .xlsx or .json export')
    if not isinstance(rows, list) or not rows or len(rows) > MAX_ROWS:
        raise ValueError(f'Expected 1–{MAX_ROWS} project records')
    if any(not isinstance(row, dict) for row in rows):
        raise ValueError('Every project must be an object')
    return rows


def needs_catalogue(rows):
    ids = {header(a) for a in ALIASES['id']}
    return any(not any(header(k) in ids and clean(v) for k, v in row.items()) for row in rows)


def canonical(row, mapping=None):
    lookup = {header(alias): key for key, label in FIELD_MAP.items() for alias in (key, label)}
    exact = {label: key for key, label in FIELD_MAP.items()}
    exact.update({key: key for key in FIELD_MAP})
    for key, aliases in ALIASES.items():
        lookup.update({header(alias): key for alias in aliases})
    lookup.update({header(label): key for key, label in (mapping or {}).items()})
    exact.update({label: key for key, label in (mapping or {}).items()})
    result = {}
    # A downloaded JSON keeps both selected source fields and their canonical names.
    source = {**(row.get('raw_fields') or {}), **row}
    for name, value in source.items():
        key = exact.get(name) or lookup.get(header(name))
        if key:
            if name in ('technology', 'research', 'themes', 'funding') and key in result and 'raw_fields' in row:
                continue  # Catalogue taxonomy arrays supplement the original export; never rewrite its source text.
            if name == 'budget_raw' and 'raw_fields' in row and key in result:
                continue  # Legacy budget_raw can be a fallback total cost, not a project budget.
            if key in result and comparable_value(key, result[key]) == comparable_value(key, value):
                continue
            result[key] = value
    if isinstance(row.get('fields'), dict):
        for key, value in row['fields'].items():
            if key in DESCRIPTIVE and (key not in result or clean(result[key]) != clean(value)):
                result[key] = value
    return result


def import_download(rows, existing, scope='spen', catalogue_rows=None, mapping=None):
    """Merge only supplied columns. Blank cells explicitly clear supplied fields."""
    old = {p['id']: p for p in existing}
    if isinstance(catalogue_rows, dict):
        catalogue_rows = catalogue_rows.get('Results', [])
    identities = {}
    for row in catalogue_rows or []:
        identities.setdefault((clean(row['OwnerNetwork']).casefold(), clean(row['ProjectReferenceNumber']).casefold()), []).append(row)
    records, errors, skipped, seen = [], [], 0, set()
    for number, row in enumerate(rows, 2):
        try:
            fields = canonical(row, mapping)
            owner = clean(fields.get('ownerNetwork'))
            if not owner:
                raise ValueError('Owner Network is required')
            if scope == 'spen' and owner not in OWNERS:
                skipped += 1
                continue
            title = clean(fields.get('projectTitle'))
            reference = clean(fields.get('projectReferenceNumber'))
            identity = clean(fields.get('id'))
            if not title:
                raise ValueError('Project Title is required')
            if not identity:
                candidates = identities.get((owner.casefold(), reference.casefold()), [])
                if not reference or len(candidates) != 1:
                    raise ValueError('Cannot resolve a unique ENA GUID from owner/reference. Use scripts.download_ena for shared or missing references')
                identity = candidates[0]['Key']
            if identity in seen:
                raise ValueError('Duplicate project GUID')
            seen.add(identity)
            record = copy.deepcopy(old.get(identity, {'id': identity, 'fields': {}, 'raw_fields': {}, 'flags': []}))
            if 'projectReferenceNumber' not in fields:
                reference = record.get('reference', '')
            record.update({'id': identity, 'title': title, 'reference': reference, 'owner': owner})
            raw = record.setdefault('raw_fields', {})
            for key, value in fields.items():
                if key in FIELD_MAP:
                    raw[FIELD_MAP[key]] = value
                    if key in raw:
                        raw[key] = value
                if key in DESCRIPTIVE:
                    text = clean(value)
                    if text:
                        record['fields'][key] = text
                    else:
                        record['fields'].pop(key, None)
            for source, target in [('projectStatus', 'status'), ('projectStartDate', 'start_date'), ('projectEndDate', 'end_date')]:
                if source in fields:
                    value = clean(fields[source])
                    if target.endswith('_date'):
                        parsed = iso_date(value)
                        if value and not parsed:
                            raise ValueError(f'Invalid {FIELD_MAP[source]}: use UK day/month/year or ISO dates')
                        value = parsed
                    record[target] = value or ('Unknown' if target == 'status' else None)
                else:
                    record.setdefault(target, 'Unknown' if target == 'status' else None)
            if record['start_date'] and record['end_date'] and record['start_date'] > record['end_date']:
                raise ValueError('Start date is after end date')
            for source, target in [('strategyTheme', 'themes'), ('technologyAreas', 'technology'), ('researchAreas', 'research'), ('fundingMechanism', 'funding')]:
                if source in fields:
                    record[target] = values(row[target] if 'raw_fields' in row and target in row else fields[source])
                else:
                    record.setdefault(target, [])
            if 'projectBudget' in fields:
                record['budget_raw'] = fields['projectBudget']
                record['budget'] = amount(fields['projectBudget'])
            record.setdefault('budget_raw', None)
            record.setdefault('budget', None)
            if 'url' in fields or not record.get('url'):
                record['url'], record['url_is_search'] = link(fields.get('url'), title)
            record['flags'] = [f for f in record.get('flags', []) if f not in ('No descriptive fields available', 'Apparent test record')]
            if not record['fields']:
                record['flags'].append('No descriptive fields available')
            if apparent_test(title):
                record['flags'].append('Apparent test record')
            # A file import is not a fresh website check. Do not accept uploaded check timestamps as proof.
            record.update({'source_checked_at': None, 'detail_last_checked_at': None,
                           'imported_at': now(), 'provenance': 'ENA file import; website check time unknown'})
            records.append(record)
        except (ValueError, TypeError, AttributeError) as exc:
            errors.append({'row': number, 'error': str(exc)})
    if not records and not errors:
        errors.append({'row': 0, 'error': 'No projects match the requested scope'})
    return records, {'rows': len(rows), 'accepted': len(records), 'outside_scope': skipped, 'errors': errors}


def download(scope='spen', progress=lambda *args: None):
    """Read ENA public rich exports; request duplicate references individually."""
    cat = catalogue()
    mapping = json.loads(request('/umbraco/api/ExportSearchResultsToCsvApi/GetProperties'))
    if not all(k in mapping for k in ('projectTitle', 'projectReferenceNumber', 'ownerNetwork')):
        raise ValueError('ENA export field schema changed')
    rows = [r for r in cat['Results'] if scope == 'all' or r['OwnerNetwork'] in OWNERS]
    if not rows:
        raise ValueError('ENA returned an empty catalogue for this scope')
    records = []
    import time
    for start in range(0, len(rows), 20):
        batch = rows[start:start + 20]
        exported, _ = export_rows(batch, mapping)
        refs = [clean(r['ProjectReferenceNumber']).casefold() for r in batch]
        for row in batch:
            ref = clean(row['ProjectReferenceNumber']).casefold()
            if not ref or refs.count(ref) > 1:
                single, _ = export_rows([row], mapping)
                records.extend(reconcile([row], single, mapping, owners=None if scope == 'all' else OWNERS))
                time.sleep(1)
            else:
                match = [e for e in exported if clean(e.get(mapping['projectReferenceNumber'])).casefold() == ref]
                records.extend(reconcile([row], match, mapping, owners=None if scope == 'all' else OWNERS))
        progress(min(start + 20, len(rows)), len(rows))
        if start + 20 < len(rows):
            time.sleep(1)
    if len(records) != len(rows) or len({r['id'] for r in records}) != len(rows):
        raise ValueError('ENA download is incomplete; no output replaced')
    return {'format': 'ena-map-download-v1', 'downloaded_at': now(), 'scope': scope,
            'projects': records, 'catalogue': cat, 'field_mapping': mapping}
